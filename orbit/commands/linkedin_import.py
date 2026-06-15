"""orbit linkedin import — ingest LinkedIn data exports into Orbit's DB.

Supports three LinkedIn-side artifacts that complement each other:

1. AggregateAnalytics_*.xlsx — what you download from LinkedIn Analytics page.
   Has the TOP POSTS sheet: 50 highest-engagement + 50 highest-impression posts
   with URL, date, engagement count, impression count. Post text is implied via
   URL slug only. Bundles likes+comments+reposts into one 'engagements' total.

2. SinglePostAnalytics_*.xlsx — per-post export from a single post's analytics
   page ('Post analytics' key-value sheet). Carries the TRUE Reactions/Comments/
   Reposts/Saves split that the LinkedIn API can't return for personal authors.
   This is the only source of real likes/comments — use it for the recent posts
   you care about. Imported via `orbit linkedin import single <files...>`.

3. Shares.csv (from https://www.linkedin.com/mypreferences/d/download-my-data)
   Has the full post text + date for every share. No engagement counts.

All three key on the IN-URL post URN (…-share-<urn>-…), NOT the filename URN
(SinglePostAnalytics filenames carry a *different* URN than the post URL). The
in-URL URN is what API-synced rows store as posts.external_id, so importing
attaches to existing posts rather than forking duplicates. Idempotent on URN.
"""
from __future__ import annotations
import csv
import re
from datetime import datetime, timezone
from pathlib import Path

import typer
from rich.console import Console
from rich.table import Table

from orbit.db import DB

app = typer.Typer(help="Import LinkedIn data exports into Orbit's DB")
console = Console()

# linkedin.com/posts/<user>_<slug>-<type>-<urn>[-suffix]
# type can be: share, ugcPost, activity
_POST_URL_RE = re.compile(
    r"/posts/[^/]+_(.*?)-(share|ugcPost|activity)-(\d+)"
)


@app.command("analytics")
def import_analytics(
    xlsx_path: Path = typer.Argument(..., exists=True, help="AggregateAnalytics_*.xlsx from LinkedIn Analytics"),
) -> None:
    """Import the TOP POSTS sheet from a LinkedIn aggregate-analytics export."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        console.print("\n[red]Need openpyxl: pip install openpyxl[/red]\n")
        raise typer.Exit(1)

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    if "TOP POSTS" not in wb.sheetnames:
        console.print(f"\n[red]No 'TOP POSTS' sheet in {xlsx_path.name}.[/red]\n")
        raise typer.Exit(1)
    ws = wb["TOP POSTS"]
    rows = list(ws.iter_rows(values_only=True))

    # Rows 0-1 are disclaimer + blank, row 2 is header, data from row 3.
    # Cols: 0=eng URL, 1=eng date, 2=eng count, 4=imp URL, 5=imp date, 6=imp count
    posts: dict[str, dict] = {}  # key = LinkedIn URN
    for r in rows[3:]:
        for url_idx, date_idx, val_idx, val_key in [
            (0, 1, 2, "engagements"), (4, 5, 6, "impressions"),
        ]:
            url = r[url_idx] if url_idx < len(r) else None
            if not url:
                continue
            parsed = _parse_post_url(url)
            if not parsed:
                continue
            slug, post_type, urn = parsed
            entry = posts.setdefault(urn, {
                "urn": urn,
                "slug": slug,
                "post_type": post_type,
                "url": url,
                "published_at": _to_iso(r[date_idx] if date_idx < len(r) else None),
            })
            entry[val_key] = int(r[val_idx] or 0) if val_idx < len(r) else 0

    console.print(f"\nParsed [bold]{len(posts)}[/bold] unique posts from {xlsx_path.name}")

    # Insert into DB. Idempotent: keyed on posts.external_id (= LinkedIn URN).
    db = DB()
    now = datetime.now(timezone.utc).isoformat()
    inserted_posts, updated_metrics = 0, 0

    for urn, p in posts.items():
        # 1. posts row (one per LinkedIn URN)
        existing = list(db.conn["posts"].rows_where("external_id = ?", [urn]))
        if existing:
            post_row_id = existing[0]["id"]
        else:
            # Need a drafts row for FK. Use slug as content placeholder until
            # the user runs `orbit linkedin import shares` to backfill text.
            idea_id = db.conn["ideas"].insert({
                "title": p["slug"][:80].replace("-", " "),
                "source": "import:linkedin-analytics",
                "status": "published",
                "created_at": p["published_at"] or now,
            }).last_pk
            draft_id = db.conn["drafts"].insert({
                "idea_id": idea_id,
                "platform": "linkedin",
                "content": p["slug"].replace("-", " "),  # placeholder until Shares.csv
                "status": "published",
                "source_url": p["url"],
                "created_at": p["published_at"] or now,
            }).last_pk
            post_row_id = db.conn["posts"].insert({
                "draft_id": draft_id,
                "platform": "linkedin",
                "external_id": urn,
                "published_at": p["published_at"] or now,
            }).last_pk
            inserted_posts += 1

        # 2. metrics row (one per (post, pulled_at))
        metric = {
            "post_id": post_row_id,
            "platform": "linkedin",
            "views": int(p.get("impressions") or 0),
            "likes": 0,         # not split out in this export
            "comments": 0,      # not split out
            "shares": int(p.get("engagements") or 0),  # repurpose 'shares' as total engagements
            "reposts": 0,
            "clicks": 0,
            "saves": 0,
            "pulled_at": now,
        }
        existing_metric = list(db.conn["metrics"].rows_where(
            "post_id = ? AND pulled_at = ?", [post_row_id, now]
        ))
        if not existing_metric:
            db.conn["metrics"].insert(metric)
            updated_metrics += 1

    console.print(f"  [green]✓[/green] {inserted_posts} new posts, {updated_metrics} metrics rows")
    console.print(
        f"\n[dim]Engagement counts stored in `metrics.shares` (LinkedIn export bundles "
        f"likes+comments+reposts into one 'engagements' total).[/dim]"
    )
    console.print(
        f"[dim]To add full post text: download Shares.csv from "
        f"https://www.linkedin.com/mypreferences/d/download-my-data, then run:[/dim]"
    )
    console.print(f"  [cyan]orbit linkedin import shares <path-to-Shares.csv>[/cyan]\n")


@app.command("shares")
def import_shares(
    csv_path: Path = typer.Argument(..., exists=True, help="Shares.csv from LinkedIn data export"),
) -> None:
    """Backfill full post text from a LinkedIn Shares.csv export.

    Joins by post URN against rows already inserted by `orbit linkedin import analytics`.
    Posts without a matching analytics row get inserted fresh (no engagement data).
    """
    db = DB()
    backfilled, inserted = 0, 0
    now = datetime.now(timezone.utc).isoformat()

    with csv_path.open(encoding="utf-8") as f:
        reader = csv.DictReader(f)
        # LinkedIn's Shares.csv columns historically: Date, ShareLink, ShareCommentary, SharedUrl, MediaUrl, Visibility
        # Column names have changed over the years; be defensive.
        for row in reader:
            url = row.get("ShareLink") or row.get("Share Link") or ""
            text = row.get("ShareCommentary") or row.get("Commentary") or row.get("Share Commentary") or ""
            date = row.get("Date") or ""
            if not url or not text:
                continue
            parsed = _parse_post_url(url)
            if not parsed:
                continue
            _slug, _ptype, urn = parsed

            # Find the drafts row via posts → draft_id (analytics import already created it)
            existing_posts = list(db.conn["posts"].rows_where("external_id = ?", [urn]))
            if existing_posts:
                # Backfill content on the corresponding draft
                draft_row = list(db.conn["drafts"].rows_where("id = ?", [existing_posts[0]["draft_id"]]))
                if draft_row:
                    db.conn["drafts"].update(draft_row[0]["id"], {"content": text})
                    backfilled += 1
            else:
                # Fresh insert (no analytics row for this post)
                pub_iso = _to_iso(date) or now
                idea_id = db.conn["ideas"].insert({
                    "title": text[:80],
                    "source": "import:linkedin-shares",
                    "status": "published",
                    "created_at": pub_iso,
                }).last_pk
                draft_id = db.conn["drafts"].insert({
                    "idea_id": idea_id,
                    "platform": "linkedin",
                    "content": text,
                    "status": "published",
                    "source_url": url,
                    "created_at": pub_iso,
                }).last_pk
                db.conn["posts"].insert({
                    "draft_id": draft_id,
                    "platform": "linkedin",
                    "external_id": urn,
                    "published_at": pub_iso,
                })
                inserted += 1

    console.print(
        f"\n[green]✓[/green] Backfilled text on {backfilled} posts, inserted {inserted} new "
        f"(no analytics) from {csv_path.name}\n"
    )


@app.command("single")
def import_single_cmd(
    xlsx_paths: list[Path] = typer.Argument(
        ..., exists=True,
        help="One or more SinglePostAnalytics_*.xlsx exports from a post's analytics page",
    ),
) -> None:
    """Import per-post SinglePostAnalytics exports (real reaction/comment breakdown).

    Unlike `import analytics` (the aggregate TOP POSTS sheet), these per-post files
    carry the true Reactions/Comments/Reposts/Saves split that the LinkedIn API
    cannot return for personal authors. Idempotent on the in-URL post URN.
    """
    import_single(xlsx_paths)


def parse_single_post_analytics(xlsx_path: Path) -> dict | None:
    """Pure parser: SinglePostAnalytics_*.xlsx -> normalized metrics dict.

    Returns None if the post URL can't be parsed into a URN. The sheet is a
    key-value layout (label in col A, value in col B); the lower demographics
    block has a 3rd column, so we only treat 2-col rows as metrics.

    Mapping: views=Impressions, likes=Reactions, comments=Comments,
    reposts=Reposts, shares=Reposts (mirror), saves=Saves, clicks=0.
    """
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["Post analytics"] if "Post analytics" in wb.sheetnames else wb[wb.sheetnames[0]]

    kv: dict[str, object] = {}
    for r in ws.iter_rows(values_only=True):
        if not r or r[0] is None or len(r) < 2 or r[1] is None:
            continue
        # Skip 3-col demographic rows ('Job title' | 'Founder' | '7%').
        if len(r) >= 3 and r[2] is not None:
            continue
        kv[str(r[0]).strip()] = r[1]

    url = str(kv.get("Post URL") or "")
    parsed = _parse_post_url(url)
    if not parsed:
        return None
    _slug, _ptype, urn = parsed

    def _int(key: str) -> int:
        try:
            return int(kv.get(key) or 0)
        except (TypeError, ValueError):
            return 0

    reposts = _int("Reposts")
    return {
        "urn": urn,
        "url": url,
        "published_at": _to_iso(kv.get("Post Date")),
        "views": _int("Impressions"),
        "likes": _int("Reactions"),
        "comments": _int("Comments"),
        "reposts": reposts,
        "shares": reposts,  # mirror — locked mapping 2026-06-15
        "saves": _int("Saves"),
        "clicks": 0,  # not present in this export
        "social_engagements": _int("Social engagements"),
    }


def import_single(xlsx_paths: list[Path], *, db: DB | None = None) -> None:
    """I/O wrapper: parse each SinglePostAnalytics file and upsert into the DB.

    Idempotent on the in-URL post URN (`posts.external_id`). Attaches a fresh
    time-stamped metrics row to the matching post, creating placeholder
    ideas/drafts/posts rows only when the post isn't already known.
    """
    db = db or DB()
    now = datetime.now(timezone.utc).isoformat()
    inserted_posts, metric_rows, skipped = 0, 0, 0

    for path in xlsx_paths:
        p = parse_single_post_analytics(Path(path))
        if not p:
            console.print(f"  [yellow]⚠ skipped[/yellow] (no URN): {Path(path).name}")
            skipped += 1
            continue

        urn = p["urn"]
        existing = list(db.conn["posts"].rows_where("external_id = ?", [urn]))
        if existing:
            post_row_id = existing[0]["id"]
        else:
            idea_id = db.conn["ideas"].insert({
                "title": (p["url"].split("/posts/")[-1][:80].replace("-", " ") or urn),
                "source": "import:linkedin-single",
                "status": "published",
                "created_at": p["published_at"] or now,
            }).last_pk
            draft_id = db.conn["drafts"].insert({
                "idea_id": idea_id,
                "platform": "linkedin",
                "content": "",  # backfill text via `orbit linkedin import shares`
                "status": "published",
                "source_url": p["url"],
                "created_at": p["published_at"] or now,
            }).last_pk
            post_row_id = db.conn["posts"].insert({
                "draft_id": draft_id,
                "platform": "linkedin",
                "external_id": urn,
                "published_at": p["published_at"] or now,
            }).last_pk
            inserted_posts += 1

        db.conn["metrics"].insert({
            "post_id": post_row_id,
            "platform": "linkedin",
            "views": p["views"],
            "likes": p["likes"],
            "comments": p["comments"],
            "shares": p["shares"],
            "reposts": p["reposts"],
            "clicks": p["clicks"],
            "saves": p["saves"],
            "pulled_at": now,
        })
        metric_rows += 1

    console.print(
        f"\n[green]✓[/green] {metric_rows} metrics rows "
        f"({inserted_posts} new posts, {skipped} skipped)\n"
    )


# --- helpers ---


def _parse_post_url(url: str | None) -> tuple[str, str, str] | None:
    """Return (slug, post_type, urn) or None.

    Handles two URL shapes LinkedIn uses for the same data:
    - Analytics export:  linkedin.com/posts/<user>_<slug>-<share|ugcPost|activity>-<urn>[-suffix]
    - Shares.csv export: linkedin.com/feed/update/urn%3Ali%3A<share|ugcPost|activity>%3A<urn>
    """
    if not url:
        return None
    # Shape 1: /posts/<user>_<slug>-<type>-<urn>
    m = _POST_URL_RE.search(url)
    if m:
        return m.group(1), m.group(2), m.group(3)
    # Shape 2: /feed/update/urn%3Ali%3A<type>%3A<urn>
    import urllib.parse
    decoded = urllib.parse.unquote(url)
    m2 = re.search(r"urn:li:(share|ugcPost|activity):(\d+)", decoded)
    if m2:
        # No slug available in this URL shape — pass empty so caller can fall back
        # to the analytics row's existing slug or to the post text itself.
        return "", m2.group(1), m2.group(2)
    return None


def _to_iso(date_val) -> str | None:
    """Coerce a LinkedIn date string ('5/27/2025') or datetime to ISO."""
    if not date_val:
        return None
    if isinstance(date_val, datetime):
        return date_val.replace(tzinfo=timezone.utc).isoformat()
    if isinstance(date_val, str):
        for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y"):
            try:
                return datetime.strptime(date_val, fmt).replace(tzinfo=timezone.utc).isoformat()
            except ValueError:
                continue
    return None
